import re
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document

from app.services.ocr_service import VisionOCRService


@dataclass
class ExtractedDocument:
    file_name: str
    text: str
    used_ocr: bool
    page_count: int


def extract_docx_text(docx_path: Path) -> str:
    document = Document(str(docx_path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    return "\n".join([text for text in paragraphs if text])


def _normalize_text(text: str) -> str:
    normalized = (text or "").replace("\u00a0", " ").replace("\t", " ")
    normalized = re.sub(r"[ ]{2,}", " ", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _is_gibberish_like(text: str) -> bool:
    compact = " ".join((text or "").split()).strip()
    if not compact:
        return True

    # A lot of isolated single-letter tokens usually indicates poor extraction.
    if re.search(r"(?:\b[a-zA-Z]\b\s+){6,}", compact):
        return True

    tokens = re.findall(r"[A-Za-z0-9\-]+", compact)
    if len(tokens) < 8:
        return True

    alpha_tokens = [token for token in tokens if re.search(r"[A-Za-z]", token)]
    if not alpha_tokens:
        return True

    single_letters = [token for token in alpha_tokens if len(token) == 1]
    if single_letters and (len(single_letters) / max(1, len(alpha_tokens))) > 0.28:
        return True

    # Very low unique ratio can indicate repeated OCR noise patterns.
    unique_ratio = len(set(tokens)) / max(1, len(tokens))
    if unique_ratio < 0.22:
        return True

    return False


def _text_quality_score(text: str) -> float:
    compact = " ".join((text or "").split()).strip()
    if not compact:
        return 0.0

    tokens = re.findall(r"[A-Za-z0-9\-]+", compact)
    if not tokens:
        return 0.0

    alpha_tokens = [token for token in tokens if re.search(r"[A-Za-z]", token)]
    alpha_ratio = len(alpha_tokens) / len(tokens)

    long_tokens = [token for token in alpha_tokens if len(token) >= 3]
    long_ratio = len(long_tokens) / max(1, len(alpha_tokens))

    single_letter_ratio = len(
        [token for token in alpha_tokens if len(token) == 1]
    ) / max(1, len(alpha_tokens))

    # Heuristic score in [0, 1]
    score = (
        (0.45 * alpha_ratio)
        + (0.45 * long_ratio)
        + (0.10 * (1.0 - min(1.0, single_letter_ratio * 2.0)))
    )
    return max(0.0, min(1.0, score))


def extract_pdf_page_texts(pdf_path: Path) -> tuple[list[str], int]:
    document = fitz.open(pdf_path)
    page_texts: list[str] = []

    try:
        for page in document:
            text = page.get_text("text")
            page_texts.append(_normalize_text(text or ""))
        page_count = document.page_count
    finally:
        document.close()

    return page_texts, page_count


def extract_pdf_text(pdf_path: Path) -> tuple[str, int]:
    page_texts, page_count = extract_pdf_page_texts(pdf_path)
    text = "\n\n".join(page for page in page_texts if page.strip())
    return text, page_count


def is_low_text_pdf(text: str, page_count: int, low_text_char_threshold: int) -> bool:
    if page_count <= 0:
        return True

    cleaned = _normalize_text(text)
    text_chars = len(cleaned)
    average_chars_per_page = text_chars / page_count

    # If extracted text is too short OR looks mostly noisy, OCR should be considered.
    return (
        text_chars < low_text_char_threshold
        or average_chars_per_page < 90
        or _is_gibberish_like(cleaned)
    )


def identify_low_text_pages(
    page_texts: list[str], min_chars_per_page: int = 90
) -> list[int]:
    low_pages: list[int] = []
    for index, text in enumerate(page_texts):
        cleaned = _normalize_text(text)
        if len(cleaned) < min_chars_per_page or _is_gibberish_like(cleaned):
            low_pages.append(index)
    return low_pages


def _should_replace_with_ocr(direct_text: str, ocr_text: str) -> bool:
    direct_clean = _normalize_text(direct_text)
    ocr_clean = _normalize_text(ocr_text)

    if not ocr_clean:
        return False
    if not direct_clean:
        return True

    direct_score = _text_quality_score(direct_clean)
    ocr_score = _text_quality_score(ocr_clean)

    # Prefer OCR when direct extraction is weak/noisy.
    if _is_gibberish_like(direct_clean) and not _is_gibberish_like(ocr_clean):
        return True

    # Prefer OCR if quality is clearly better or direct text is too short.
    if len(direct_clean) < 120 and len(ocr_clean) >= 60:
        return True
    if (ocr_score - direct_score) >= 0.12:
        return True
    if len(ocr_clean) > (len(direct_clean) * 1.35):
        return True

    return False


def _merge_direct_and_ocr_pages(
    page_texts: list[str],
    ocr_page_map: dict[int, str],
) -> str:
    merged_pages: list[str] = []

    for index, direct_text in enumerate(page_texts):
        direct_clean = _normalize_text(direct_text)
        ocr_clean = _normalize_text(ocr_page_map.get(index, ""))

        if _should_replace_with_ocr(direct_clean, ocr_clean):
            merged_pages.append(ocr_clean)
        else:
            merged_pages.append(direct_clean or ocr_clean)

    return "\n\n".join(page for page in merged_pages if page.strip())


def extract_file_text(
    file_path: Path,
    low_text_char_threshold: int,
    ocr_service: VisionOCRService | None = None,
) -> ExtractedDocument:
    extension = file_path.suffix.lower()

    if extension == ".docx":
        text = extract_docx_text(file_path)
        return ExtractedDocument(
            file_name=file_path.name,
            text=_normalize_text(text),
            used_ocr=False,
            page_count=1,
        )

    if extension == ".pdf":
        page_texts, page_count = extract_pdf_page_texts(file_path)
        direct_text = "\n\n".join(page for page in page_texts if page.strip())
        direct_text = _normalize_text(direct_text)

        if is_low_text_pdf(direct_text, page_count, low_text_char_threshold):
            if ocr_service is None:
                return ExtractedDocument(
                    file_name=file_path.name,
                    text=direct_text,
                    used_ocr=False,
                    page_count=page_count,
                )

            low_text_pages = identify_low_text_pages(page_texts)
            pages_to_ocr = low_text_pages if low_text_pages else list(range(page_count))

            ocr_page_map = ocr_service.ocr_pdf_pages(
                pdf_path=file_path,
                page_numbers=pages_to_ocr,
            )

            merged_text = _merge_direct_and_ocr_pages(
                page_texts=page_texts, ocr_page_map=ocr_page_map
            )

            # If merged text is still weak and OCR was partial, run full OCR once.
            if (
                ocr_page_map
                and _is_gibberish_like(merged_text)
                and len(pages_to_ocr) < max(1, page_count)
            ):
                full_ocr_page_map = ocr_service.ocr_pdf_pages(
                    pdf_path=file_path,
                    page_numbers=list(range(page_count)),
                )
                merged_text = _merge_direct_and_ocr_pages(
                    page_texts=page_texts,
                    ocr_page_map=full_ocr_page_map,
                )
                ocr_page_map = full_ocr_page_map

            return ExtractedDocument(
                file_name=file_path.name,
                text=_normalize_text(merged_text),
                used_ocr=bool(ocr_page_map),
                page_count=page_count,
            )

        return ExtractedDocument(
            file_name=file_path.name,
            text=direct_text,
            used_ocr=False,
            page_count=page_count,
        )

    raise ValueError(f"Unsupported file type for extraction: {extension}")
